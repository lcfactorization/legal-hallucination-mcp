"""H-7/H-8 新维度测试 — 裁判结果-援引依据一致性 & 诉请-判决回应一致性。

v0.4.0 新增：
- 维度数 6→8 断言；
- 权重和 = 1.00 断言；
- H-7 "主文判项无说理呼应" 合成样例 → critical flag；
- H-8 "诉求未回应 / 变更未审理" 合成样例 → flag；
- 法规库 .docx 加载：无 python-docx 时优雅跳过（monkeypatch 模拟 ImportError）。

脱敏说明：本文书样例均为合成数据，当事人以"某劳动者/某用人单位"占位，
案号以（20XX）X民终XXXX号占位，不含任何真实人名、真实公司名、真实案号。
"""
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

import pytest

from legal_hallucination_mcp.config import (
    DIMENSION_ORDER,
    DIMENSION_TITLES,
    DIMENSION_WEIGHTS,
    HALLUCINATION_DIMENSIONS,
)
from legal_hallucination_mcp.law_knowledge_base import LawKnowledgeBase
from legal_hallucination_mcp.report_builder import ReportBuilder
from legal_hallucination_mcp.rule_engine import RuleEngine

H7_UNSUPPORTED = """某市人民法院刑事附带民事判决书

某劳动者诉某用人单位：请求支付拖欠工资50,000元。

本院认为：本案事实清楚，证据充分。

判决如下：
一、某用人单位于本判决生效之日起十日内支付某劳动者工资50,000元；
"""

H7_SUPPORTED = """某市人民法院民事判决书

某劳动者诉某用人单位：请求支付拖欠工资50,000元。

本院认为：根据《某工资支付条例》第10条的规定，用人单位应当按时足额支付工资。
本案中，某用人单位拖欠某劳动者工资（见《证据1.md》），故应予支持。

判决如下：
一、某用人单位于本判决生效之日起十日内支付某劳动者工资50,000元；
"""

H8_CLAIM_UNANSWERED = """某市人民法院民事判决书

原告诉称：1. 判令某用人单位支付拖欠工资50,000元；
2. 判令某用人单位支付解除劳动合同经济补偿金30,000元。

本院认为：某用人单位拖欠某劳动者工资的事实清楚（见《证据1.md》）。

判决如下：
一、某用人单位于本判决生效之日起十日内支付某劳动者工资50,000元；
"""

H8_AMENDED_UNTRIED = """某市人民法院民事判决书

原告诉称：1. 判令某用人单位支付拖欠工资50,000元。
庭审中，某劳动者变更诉讼请求，将拖欠工资请求增加至80,000元。

本院认为：某用人单位拖欠某劳动者工资的事实清楚（见《证据1.md》）。

判决如下：
一、某用人单位于本判决生效之日起十日内支付某劳动者工资50,000元；
"""


class TestDimensionMetadata:
    def test_dimensions_count_is_8(self):
        """维度数断言更新：v0.3.x 为 6 维，v0.4.0 起为 8 维（H-1~H-8）。"""
        assert len(HALLUCINATION_DIMENSIONS) == 8
        assert len(DIMENSION_TITLES) == 8
        assert len(DIMENSION_ORDER) == 8
        assert len(DIMENSION_WEIGHTS) == 8

    def test_new_dimension_keys_present(self):
        assert "h7_judgment_basis_consistency" in HALLUCINATION_DIMENSIONS
        assert "h8_claim_response_consistency" in HALLUCINATION_DIMENSIONS
        assert DIMENSION_ORDER["h7_judgment_basis_consistency"] == 7
        assert DIMENSION_ORDER["h8_claim_response_consistency"] == 8

    def test_dimension_weights_sum_to_one(self):
        assert abs(sum(DIMENSION_WEIGHTS.values()) - 1.00) < 1e-9, \
            f"八维权重之和应为1.00，实际为{sum(DIMENSION_WEIGHTS.values())}"

    def test_dimension_weights_values(self):
        expected = {
            "h1_sourceless_fabrication": 0.22,
            "h2_law_misapplication": 0.18,
            "h3_evidence_chain_break": 0.18,
            "h4_subjective_rhetoric": 0.08,
            "h5_claim_boundary_breach": 0.15,
            "h6_nontext_evidence_fail": 0.04,
            "h7_judgment_basis_consistency": 0.075,
            "h8_claim_response_consistency": 0.075,
        }
        assert DIMENSION_WEIGHTS == expected


class TestH7JudgmentBasisConsistency:
    def test_unsupported_item_flags_critical(self):
        engine = RuleEngine()
        flags = engine.check_judgment_basis_consistency(H7_UNSUPPORTED)
        assert len(flags) > 0, "判项无说理呼应时应产出H-7标记"
        critical = [f for f in flags if f.severity == "critical"]
        assert len(critical) > 0, "H-7 无依据支撑判项应为critical级"
        assert all(f.h_code == "H-7" for f in flags)
        assert all(f.sub_type == "裁判结果缺乏依据支撑" for f in flags)

    def test_supported_item_no_flag(self):
        engine = RuleEngine()
        flags = engine.check_judgment_basis_consistency(H7_SUPPORTED)
        assert len(flags) == 0, "说理含法条引用与关键词呼应的判项不应被标记"

    def test_flags_enter_dimension_and_score(self):
        engine = RuleEngine()
        result = engine.run_full_scan(document_text=H7_UNSUPPORTED)
        h7_dims = [d for d in result.dimensions if d.h_code == "H-7"]
        assert len(h7_dims) == 1, "结果应包含H-7维度"
        assert h7_dims[0].critical_count > 0
        assert result.hallucination_score > 0, "H-7标记应进入罚分路径"
        assert result.risk_grade in ("C", "D", "F"), "critical级H-7应抬高风险等级"


class TestH8ClaimResponseConsistency:
    def test_unanswered_claim_flags_high(self):
        engine = RuleEngine()
        flags = engine.check_claim_response_consistency(H8_CLAIM_UNANSWERED, "")
        not_addressed = [f for f in flags if f.sub_type == "诉讼请求未获回应"]
        assert len(not_addressed) > 0, "诉请未获回应时应产出H-8标记"
        assert any(f.severity == "high" for f in not_addressed), "未回应诉请应为high级"
        assert all(f.h_code == "H-8" for f in not_addressed)

    def test_amended_claim_untried_flags_critical(self):
        engine = RuleEngine()
        flags = engine.check_claim_response_consistency(H8_AMENDED_UNTRIED, "")
        amended = [f for f in flags if f.sub_type == "变更后诉讼请求未审理"]
        assert len(amended) == 1, "变更/增加诉请未审理应产出H-8 critical标记"
        assert amended[0].severity == "critical"
        assert amended[0].h_code == "H-8"

    def test_flags_enter_dimension_and_score(self):
        engine = RuleEngine()
        result = engine.run_full_scan(document_text=H8_CLAIM_UNANSWERED)
        h8_dims = [d for d in result.dimensions if d.h_code == "H-8"]
        assert len(h8_dims) == 1
        assert h8_dims[0].high_count > 0
        assert result.hallucination_score > 0


class TestScorePenaltyPath:
    def test_severity_coefficients_via_server_tool(self):
        """H-7/H-8 标志应可进入 calculate_hallucination_score 的语义罚分路径
        （severity 系数口径：critical 5 / high 3 / medium 1.5 / low 0.5）。"""
        import json

        from legal_hallucination_mcp.server import calculate_hallucination_score

        engine = RuleEngine()
        result = engine.run_full_scan(document_text=H7_UNSUPPORTED)
        rule_json = json.dumps({
            "hallucination_score": result.hallucination_score,
            "risk_grade": result.risk_grade,
        })
        semantic = [{"flags": [{"severity": "critical"}, {"severity": "high"}]}]
        out = json.loads(calculate_hallucination_score(
            rule_results_json=rule_json,
            semantic_results_json=json.dumps(semantic),
        ))
        assert out["success"] is True
        assert out["semantic_penalty"] == pytest.approx(5.0 + 3.0)
        assert out["total_score"] == pytest.approx(round(min(result.hallucination_score + 8.0, 100.0), 1))

    def test_run_rule_engine_full_exposes_h7_h8_flags(self):
        import json

        from legal_hallucination_mcp.server import run_rule_engine_full

        out = json.loads(run_rule_engine_full(document_text=H8_CLAIM_UNANSWERED))
        assert out["success"] is True
        assert any(d["h_code"] == "H-8" for d in out["dimensions"])
        assert len(out["h7_h8_flags"]) > 0, "run_rule_engine_full应导出H-7/H-8标志供语义罚分路径使用"
        assert all(f["h_code"] in ("H-7", "H-8") for f in out["h7_h8_flags"])


class TestReportSections:
    def test_report_contains_h7_h8_sections(self):
        engine = RuleEngine()
        result = engine.run_full_scan(document_text=H8_CLAIM_UNANSWERED)
        report = ReportBuilder().build_report(result)
        assert "H-7：裁判结果—援引依据一致性检测" in report
        assert "H-8：诉请—判决回应一致性检测" in report
        assert "诉讼请求未获回应" in report, "H-8未回应诉请应以幻觉卡片形式呈现"

    def test_report_section_numbering_ascending(self):
        engine = RuleEngine()
        result = engine.run_full_scan(document_text=H7_UNSUPPORTED)
        report = ReportBuilder().build_report(result)
        markers = ["七、H-1", "八、H-2", "九、H-3", "十、H-4", "十一、H-5",
                   "十二、H-6", "十三、H-7", "十四、H-8"]
        positions = [report.find(m) for m in markers]
        assert all(p >= 0 for p in positions), f"报告应包含全部维度小节：{positions}"
        assert positions == sorted(positions), "维度小节应按 H-1→H-8 顺序排列"


class TestDocxKnowledgeBaseLoading:
    def _make_docx(self, path):
        from docx import Document

        doc = Document()
        doc.add_paragraph("《某测试条例》")
        doc.add_paragraph("第一条 用人单位应当按时足额支付工资。")
        doc.save(str(path))

    def test_docx_loaded_when_python_docx_available(self, tmp_path):
        pytest.importorskip("docx")
        self._make_docx(tmp_path / "某测试条例.docx")
        (tmp_path / "某说明.txt").write_text("辅助说明，无法条", encoding="utf-8")

        kb = LawKnowledgeBase()
        res = kb.load_from_directory(str(tmp_path))
        assert res["success"] is True
        assert res["files"] == 2
        assert res["articles"] >= 1, "应从.docx按段落提取法条"

    def test_docx_skipped_gracefully_without_python_docx(self, tmp_path, monkeypatch):
        """无 python-docx 时应优雅跳过 .docx 并汇总 warning，不抛异常、不加硬依赖。"""
        self._make_docx(tmp_path / "某测试条例.docx")
        (tmp_path / "某条例.md").write_text(
            "《某条例》\n第一条 用人单位应当按时足额支付工资。\n", encoding="utf-8",
        )

        # 模拟 ImportError：将 docx 模块置为 None，`from docx import Document` 将抛 ImportError
        monkeypatch.setitem(sys.modules, "docx", None)

        kb = LawKnowledgeBase()
        res = kb.load_from_directory(str(tmp_path))
        assert res["success"] is True
        assert res["files"] == 2, ".docx文件应被计数"
        assert res["articles"] >= 1, ".md文件应正常解析"
        assert kb.loaded is True
